# -*- coding: utf-8 -*-
"""New deep-dive deck for churn defense extras."""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUT = ROOT / "Churn_deep_pack.pptx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
LINE = RGBColor(0xE6, 0xE6, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W = Inches(13.333)
H = Inches(7.5)
ML = Inches(0.72)
MR = Inches(0.72)


def _run(p, text, size=18, bold=False, color=INK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return r


def tf_box(slide, l, t, w, h):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def write(tf, lines):
    first = True
    for item in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(item.get("after", 6))
        _run(p, item["text"], item.get("size", 16), item.get("bold", False), item.get("color", INK))


def hline(slide, y):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ML, y, W - ML - MR, Emu(12700))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LINE
    sh.line.fill.background()


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    spTree = s.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return s


def _col_weights(rows):
    n = len(rows[0])
    if n == 2:
        return [1.2, 1.8]
    if n == 3:
        return [1.2, 1.2, 1.6]
    if n == 4:
        return [1.1, 1.1, 0.9, 1.4]
    if n == 5:
        return [1.0] * 5
    return [1] * n


def add_table(slide, left, top, width, height, rows, font_size=13):
    from lxml import etree
    from pptx.oxml.ns import qn

    n_rows, n_cols = len(rows), len(rows[0])
    ts = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = ts.table
    weights = _col_weights(rows)
    widths = [int(width * w / sum(weights)) for w in weights]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = str(rows[r][c])
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = "Calibri"
                    run.font.bold = r == 0
                    run.font.color.rgb = INK
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            lnL = etree.SubElement(tcPr, qn("a:lnL"))
            lnL.set("w", "6350")
            sf = etree.SubElement(lnL, qn("a:solidFill"))
            srgb = etree.SubElement(sf, qn("a:srgbClr"))
            srgb.set("val", "E6E6E6")
    return table


def heading(s, kicker, title):
    tf = tf_box(s, ML, Inches(0.42), W - ML - MR, Inches(1.15))
    write(
        tf,
        [
            {"text": kicker, "size": 11, "color": MUTED, "after": 6},
            {"text": title, "size": 28, "bold": True, "after": 0},
        ],
    )
    hline(s, Inches(1.22))


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cw = W - ML - MR

    h = pd.read_csv(DATA / "deep_horizons_auc.csv")
    t = pd.read_csv(DATA / "deep_tariff_segments.csv")
    o = pd.read_csv(DATA / "deep_onboarding.csv")
    s = pd.read_csv(DATA / "deep_feature_stability.csv")
    loc = pd.read_csv(DATA / "deep_local_shaplike.csv")

    # cover
    slide = new_slide(prs)
    tf = tf_box(slide, ML, Inches(2.0), cw, Inches(4))
    write(
        tf,
        [
            {"text": "CHURN · DEEP PACK", "size": 12, "color": MUTED, "after": 14},
            {"text": "Углублённая аналитика к защите", "size": 34, "bold": True, "after": 14},
            {
                "text": "Горизонты AUC · кто пострадал после тарифа · onboarding · SHAP · playbook · стабильность фич",
                "size": 16,
                "color": MUTED,
                "after": 0,
            },
        ],
    )

    # 1 horizons AUC
    slide = new_slide(prs)
    heading(slide, "ТАРГЕТ", "Горизонты 60 / 90 / 120 — не только rate, но и AUC")
    rows = [["Горизонт", "train rate", "test rate", "AUC", "AP"]]
    for _, r in h.iterrows():
        note = "" if r["window_closed"] else "*"
        rows.append(
            [
                f"{int(r.horizon)}д{note}",
                f"{100*r.train_rate:.2f}%",
                f"{100*r.test_rate:.2f}%",
                f"{r.auc:.3f}",
                f"{r.ap:.3f}",
            ]
        )
    add_table(slide, ML, Inches(1.45), cw, Inches(3.2), rows, 14)
    tf = tf_box(slide, ML, Inches(5.0), cw, Inches(1.8))
    write(
        tf,
        [
            {
                "text": "60д: выше rate и чуть ниже AUC — больше шума временных пауз. 90д: лучший баланс rate/качества.",
                "size": 15,
                "after": 6,
            },
            {
                "text": "*120д на test 25.09 окно не полностью закрыто к 31.12 — AUC ориентировочный. Рабочий выбор: 90 дней.",
                "size": 14,
                "color": MUTED,
                "after": 0,
            },
        ],
    )

    # 2 tariff who hurt
    slide = new_slide(prs)
    heading(slide, "ТАРИФ", "После 1 июля: кто пострадал")
    # pivot-ish table
    june = t[t.snap.astype(str).str.startswith("2025-06")].set_index("grp")["churn_pct"]
    sep = t[t.snap.astype(str).str.startswith("2025-09")].set_index("grp")["churn_pct"]
    groups = ["payroll_1", "payroll_0", "mass", "affluent", "premium"]
    labels = {
        "payroll_1": "зарплатные",
        "payroll_0": "не зарплатные",
        "mass": "mass",
        "affluent": "affluent",
        "premium": "premium",
    }
    rows = [["Сегмент", "июнь 2025", "сентябрь 2025", "Δ п.п."]]
    for g in groups:
        if g in june.index and g in sep.index:
            d = float(sep[g]) - float(june[g])
            rows.append([labels[g], f"{june[g]:.2f}%", f"{sep[g]:.2f}%", f"+{d:.2f}" if d >= 0 else f"{d:.2f}"])
    add_table(slide, ML, Inches(1.45), cw, Inches(3.6), rows, 14)
    tf = tf_box(slide, ML, Inches(5.4), cw, Inches(1.5))
    write(
        tf,
        [
            {
                "text": "Сильнее всего вырос churn у не-зарплатных (+3.4 п.п.). Зарплатные держатся лучше, но тоже выросли.",
                "size": 15,
                "after": 6,
            },
            {
                "text": "Affluent/premium/mass — все сегменты ухудшились после тарифа: удар широкий, не точечный.",
                "size": 14,
                "color": MUTED,
                "after": 0,
            },
        ],
    )

    # 3 onboarding
    slide = new_slide(prs)
    heading(slide, "ПОПУЛЯЦИЯ", "Onboarding tenure < 60 вынесли отдельно")
    rows = [["Корзина", "клиентов", "alive-like", "churn %"]]
    for _, r in o.iterrows():
        rows.append([r.bucket, f"{int(r.n_customers):,}", f"{int(r.n_alive_like):,}", f"{r.churn_pct:.2f}%"])
    add_table(slide, ML, Inches(1.5), cw, Inches(2.6), rows, 15)
    tf = tf_box(slide, ML, Inches(4.5), cw, Inches(2.3))
    write(
        tf,
        [
            {
                "text": "У новичков (<60 дней) мало истории для ритма/gaps — модель main-portfolio на них не учим.",
                "size": 16,
                "after": 8,
            },
            {
                "text": "На 30.06: ~3k onboarding, churn среди «живых» ~11.7% vs ~6.4% у tenure≥60. Отдельный onboarding-churn — следующий этап.",
                "size": 15,
                "color": MUTED,
                "after": 0,
            },
        ],
    )

    # 4 feature stability
    slide = new_slide(prs)
    heading(slide, "МОДЕЛЬ", "Топ фич устойчив: 2024-12 и 2025-06")
    w1 = s[s.window.str.contains("2024-12")].head(6)
    w2 = s[s.window.str.contains("2025-06")].head(6)
    rows = [["#", "окно → 2024-12", "ΔAP", "окно → 2025-06", "ΔAP"]]
    for i in range(6):
        a = w1.iloc[i] if i < len(w1) else None
        b = w2.iloc[i] if i < len(w2) else None
        rows.append(
            [
                str(i + 1),
                a.feature if a is not None else "",
                f"{a.importance:.3f}" if a is not None else "",
                b.feature if b is not None else "",
                f"{b.importance:.3f}" if b is not None else "",
            ]
        )
    add_table(slide, ML, Inches(1.45), cw, Inches(4.0), rows, 13)
    tf = tf_box(slide, ML, Inches(5.7), cw, Inches(1.2))
    write(
        tf,
        [
            {
                "text": "Общие ядра: recency / txn_cnt / app_events / retention_failed. Не один удачный месяц — сигналы повторяются.",
                "size": 15,
                "after": 0,
            }
        ],
    )

    # 5 SHAP two clients
    slide = new_slide(prs)
    heading(slide, "SHAP", "Разбор 2 клиентов (LinearExplainer)")
    cases = list(loc.groupby("case", sort=False))
    # left and right tables
    if cases:
        case_a, df_a = cases[0]
        rows_a = [["фича", "вклад"]] + [[r.feature, f"{r.contribution:+.2f}"] for _, r in df_a.head(6).iterrows()]
        add_table(slide, ML, Inches(1.55), Inches(5.7), Inches(3.6), rows_a, 12)
        meta_a = df_a.iloc[0]
        tf = tf_box(slide, ML, Inches(5.4), Inches(5.7), Inches(1.5))
        write(
            tf,
            [
                {"text": f"Клиент {int(meta_a.customer_id)} · p={meta_a.p_churn} · target={int(meta_a.target)}", "size": 13, "bold": True, "after": 4},
                {"text": "A: high risk + fee_pain → FEE_WAIVER", "size": 13, "color": MUTED, "after": 0},
            ],
        )
    if len(cases) > 1:
        case_b, df_b = cases[1]
        rows_b = [["фича", "вклад"]] + [[r.feature, f"{r.contribution:+.2f}"] for _, r in df_b.head(6).iterrows()]
        add_table(slide, Inches(7.2), Inches(1.55), Inches(5.5), Inches(3.6), rows_b, 12)
        meta_b = df_b.iloc[0]
        tf = tf_box(slide, Inches(7.2), Inches(5.4), Inches(5.5), Inches(1.5))
        write(
            tf,
            [
                {"text": f"Клиент {int(meta_b.customer_id)} · p={meta_b.p_churn} · target={int(meta_b.target)}", "size": 13, "bold": True, "after": 4},
                {"text": "B: mono_fade + grocery → cashback / персональный оффер", "size": 13, "color": MUTED, "after": 0},
            ],
        )

    # 6 playbook
    slide = new_slide(prs)
    heading(slide, "ДЕЙСТВИЯ", "Playbook: сигнал → оффер")
    add_table(
        slide,
        ML,
        Inches(1.5),
        cw,
        Inches(4.2),
        [
            ["Сигнал", "Что видим", "Действие"],
            ["высокий p + fee_pain", "боль от комиссии / жалобы", "FEE_WAIVER / скидка на комиссию"],
            ["mono_fade + grocery silence", "лестница вниз, быт ушёл с карты", "cashback на grocery / персональный оффер"],
            ["retention_failed", "уже писали — не помогло", "не тот же канал: другой оффер / звонок"],
            ["высокий p × value", "риск дорогой маржи", "приоритет в top-K слотах"],
        ],
        14,
    )
    tf = tf_box(slide, ML, Inches(6.0), cw, Inches(1.0))
    write(
        tf,
        [
            {
                "text": "Не один оффер на всех. Сегментируем удержание по причине риска — иначе повторим бесполезный winback.",
                "size": 15,
                "color": MUTED,
                "after": 0,
            }
        ],
    )

    # 7 closing
    slide = new_slide(prs)
    heading(slide, "ИТОГ", "Что добавил deep pack")
    tf = tf_box(slide, ML, Inches(1.6), cw, Inches(5))
    write(
        tf,
        [
            {"text": "1. 90 дней подтвердили и rate, и AUC/AP — не только «красивый порог».", "size": 18, "after": 10},
            {"text": "2. После тарифа удар широкий; не-зарплатные пострадали сильнее.", "size": 18, "after": 10},
            {"text": "3. Onboarding <60 вынесен: другая задача, другой churn.", "size": 18, "after": 10},
            {"text": "4. Топ фич повторяется на разных окнах — не разовая удача.", "size": 18, "after": 10},
            {"text": "5. SHAP + playbook: от сигнала к конкретному действию.", "size": 18, "after": 0},
        ],
    )

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
